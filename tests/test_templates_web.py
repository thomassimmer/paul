from __future__ import annotations

import asyncio
from io import BytesIO

from docx import Document

from app.config import Settings
from app.models import TemplateBlueprint
from app.templates_engine import analyze, store
from app.templates_engine.default import build as build_default
from app.templates_engine.extract import extract_blocks


def _uploaded() -> bytes:
    document = Document()
    paragraph = document.add_paragraph("Camille Moreau")
    paragraph.runs[0].bold = True
    document.add_paragraph("camille@example.com")
    document.add_paragraph("Experience").runs[0].bold = True
    document.add_paragraph("A result", style="List Bullet")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _seed_custom(kind: str = "cv") -> None:
    data = _uploaded()
    store.save_custom(kind, data, store.blueprint_from_blocks(kind, extract_blocks(data)))


# --- the pages ----------------------------------------------------------------


def test_the_settings_page_offers_both_kinds(client):
    response = client.get("/settings")
    assert response.status_code == 200
    assert response.text.count('id="templates"') == 1
    assert "CV" in response.text
    assert "Cover letter" in response.text
    assert "The default template is in use" in response.text


def test_the_settings_page_warns_without_a_model(client):
    assert "guessed from the layout only" in client.get("/settings").text


def test_the_old_templates_page_redirects_to_the_settings(client):
    response = client.get("/templates", follow_redirects=False)
    assert response.status_code == 303
    assert response.headers["location"] == "/settings#templates"


def test_uploading_a_template_stores_it(client, monkeypatch):
    async def fake_analyze(settings, docx_bytes, kind):
        blueprint: TemplateBlueprint = store.blueprint_from_blocks(kind, extract_blocks(docx_bytes))
        blueprint.notes = ["The model corrected 1 of the 4 blocks."]
        return blueprint

    monkeypatch.setattr("app.templates_engine.analyze.analyze", fake_analyze)

    response = client.post(
        "/templates/cv",
        files={"file": ("my-cv.docx", _uploaded(), "application/octet-stream")},
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "template imported" in response.text
    assert "The model corrected 1 of the 4 blocks." in response.text
    assert store.has_custom("cv") is True
    assert "blocks read from your file" in client.get("/settings").text

    store.delete_custom("cv")


def test_uploading_a_non_docx_is_refused(client):
    response = client.post(
        "/templates/cv",
        files={"file": ("cv.pdf", b"not a docx", "application/pdf")},
        follow_redirects=True,
    )
    assert "Could not read this file as a .docx" in response.text
    assert store.has_custom("cv") is False


def test_uploading_without_a_file_is_refused(client):
    response = client.post("/templates/cv", data={}, follow_redirects=True)
    assert "Choose a .docx file first" in response.text


def test_the_preview_shows_the_roles_and_saves_corrections(client):
    _seed_custom()
    try:
        page = client.get("/templates/cv")
        assert page.status_code == 200
        assert "Camille Moreau" in page.text
        assert 'name="role.0"' in page.text

        response = client.post(
            "/templates/cv/roles",
            data={
                "block_count": "4",
                "role.0": "headline",
                "role.1": "contact",
                "role.2": "section_title",
                "role.3": "bullet",
            },
            follow_redirects=True,
        )

        assert "Roles saved" in response.text
        assert store.load_blueprint("cv").roles() == [
            "headline",
            "contact",
            "section_title",
            "bullet",
        ]
    finally:
        store.delete_custom("cv")


def test_the_preview_needs_an_imported_template(client):
    assert "No CV template imported yet" in client.get("/templates/cv", follow_redirects=True).text


def test_resetting_goes_back_to_the_default(client):
    _seed_custom()
    try:
        response = client.post("/templates/cv/reset", follow_redirects=True)
        assert "the default is used again" in response.text
        assert store.has_custom("cv") is False
    finally:
        store.delete_custom("cv")


def test_an_unknown_kind_is_refused(client):
    assert "Unknown template kind" in client.get("/templates/nope", follow_redirects=True).text


def test_the_settings_page_links_to_the_detected_roles(client):
    _seed_custom()
    try:
        page = client.get("/settings").text
        assert 'href="/templates/cv"' in page
    finally:
        store.delete_custom("cv")


# --- the analysis -------------------------------------------------------------


def test_analyze_falls_back_to_guesses_without_a_model():
    blueprint = asyncio.run(analyze.analyze(Settings(), _uploaded(), "cv"))
    assert blueprint.roles()[0] == "name"
    assert "No model configured" in " ".join(blueprint.notes)


def test_analyze_uses_the_model_when_there_is_one(monkeypatch):
    async def fake_complete(settings, *, schema, content, system=None, **kwargs):
        assert "Allowed roles" in content
        assert system  # the prompt file was loaded
        return schema(
            blocks=[{"index": 0, "role": "headline"}, {"index": 1, "role": "contact"}]
        )

    monkeypatch.setattr("app.templates_engine.analyze.complete_structured", fake_complete)

    blueprint = asyncio.run(analyze.analyze(Settings(model="openai/gpt-4o"), _uploaded(), "cv"))

    assert blueprint.roles()[0] == "headline"
    assert "corrected" in " ".join(blueprint.notes)


def test_analyze_survives_a_model_failure(monkeypatch):
    from app.llm import LLMError

    async def failing(settings, *, schema, content, system=None, **kwargs):
        raise LLMError("provider is down")

    monkeypatch.setattr("app.templates_engine.analyze.complete_structured", failing)

    blueprint = asyncio.run(analyze.analyze(Settings(model="openai/gpt-4o"), _uploaded(), "cv"))

    assert blueprint.roles()[0] == "name"  # the guesses stood
    assert "provider is down" in " ".join(blueprint.notes)


def test_the_default_build_is_deterministic():
    assert build_default("cv")[0] == build_default("cv")[0]
