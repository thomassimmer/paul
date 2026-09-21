from __future__ import annotations

import io

import pytest
from docx import Document

from app.profiler import cv

LONG_ENOUGH = (
    "Camille Moreau, Senior Backend Engineer. Ten years building data platforms "
    "in Rust, Python and Kafka."
)


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        created = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                created.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(text: str) -> bytes:
    """A minimal single-page PDF with a real text object and a valid xref."""
    content = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 200] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_extracts_docx_paragraphs_and_tables():
    data = _docx_bytes(
        ["Camille Moreau", "Senior Backend Engineer"],
        table=[["2019 - 2022", "Northwind Software, Backend Engineer"]],
    )
    text = cv.extract_text("cv.docx", data)
    assert "Camille Moreau" in text
    assert "Northwind Software, Backend Engineer" in text


def test_extracts_pdf_text():
    text = cv.extract_text("cv.pdf", _pdf_bytes(LONG_ENOUGH))
    assert "Camille Moreau" in text


def test_extracts_plain_text():
    text = cv.extract_text("cv.txt", LONG_ENOUGH.encode("utf-8"))
    assert text.startswith("Camille Moreau")


def test_rejects_unsupported_extension():
    with pytest.raises(cv.CvError, match="Unsupported file type"):
        cv.extract_text("cv.doc", LONG_ENOUGH.encode("utf-8"))


def test_rejects_scanned_or_empty_file():
    with pytest.raises(cv.CvError, match="Almost no text"):
        cv.extract_text("cv.txt", b"short")


def test_rejects_corrupt_docx():
    with pytest.raises(cv.CvError, match="Could not read the DOCX"):
        cv.extract_text("cv.docx", b"not a zip archive at all")


def test_rejects_corrupt_pdf():
    with pytest.raises(cv.CvError, match="Could not read the PDF"):
        cv.extract_text("cv.pdf", b"not a pdf")
