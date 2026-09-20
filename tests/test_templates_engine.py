from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.models import DraftLine, TemplateBlock, TemplateBlueprint
from app.templates_engine import default, extract, pdf, render, roles, store


def _docx_bytes(*paragraphs: tuple[str, dict]) -> bytes:
    document = Document()
    for text, options in paragraphs:
        paragraph = document.add_paragraph(style=options.get("style"))
        run = paragraph.add_run(text)
        run.bold = options.get("bold", False)
        size = options.get("size")
        if size:
            run.font.size = Pt(size)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    """python-docx has no public API for this: build the element by hand."""
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    link.append(run)
    paragraph._p.append(link)


def _linked_docx() -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Reach me: ")
    _add_hyperlink(paragraph, "camille@example.com", "mailto:camille@example.com")
    paragraph.add_run(" · github.com/camille")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- extraction ---------------------------------------------------------------


def test_extracts_non_empty_paragraphs_with_a_hint():
    blocks = extract.extract_blocks(
        _docx_bytes(
            ("Camille Moreau", {"bold": True, "size": 22}),
            ("", {}),
            ("Experience", {"bold": True}),
        )
    )
    assert [block.text for block in blocks] == ["Camille Moreau", "Experience"]
    assert "bold" in blocks[0].hint
    assert "22pt" in blocks[0].hint
    assert blocks[0].xml.startswith("<w:p")


def test_reads_paragraphs_inside_tables():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "2022 - 2024"
    table.cell(0, 1).text = "Acme — Lead Engineer"
    buffer = BytesIO()
    document.save(buffer)

    blocks = extract.extract_blocks(buffer.getvalue())
    assert [block.text for block in blocks] == ["2022 - 2024", "Acme — Lead Engineer"]


def test_refuses_a_file_that_is_not_a_docx():
    with pytest.raises(extract.TemplateError, match="Could not read this file"):
        extract.extract_blocks(b"not a zip")


def test_refuses_an_empty_document():
    with pytest.raises(extract.TemplateError, match="no text"):
        extract.extract_blocks(_docx_bytes(("", {})))


# --- roles --------------------------------------------------------------------


def test_guesses_roles_from_the_layout():
    blocks = [
        extract.ExtractedBlock("Camille Moreau", "<w:p/>", "bold, 22pt"),
        extract.ExtractedBlock("camille@example.com", "<w:p/>", ""),
        extract.ExtractedBlock("Experience", "<w:p/>", "bold"),
        extract.ExtractedBlock("A result", "<w:p/>", "List Bullet, list"),
        extract.ExtractedBlock("2022 - 2024", "<w:p/>", "bold"),
    ]
    assert roles.guess_roles(
        [b.text for b in blocks], [b.hint for b in blocks], "cv"
    ) == ["name", "contact", "section_title", "bullet", "entry_dates"]


def test_guesses_the_letter_shapes():
    blocks = [
        extract.ExtractedBlock("Camille Moreau", "<w:p/>", "bold"),
        extract.ExtractedBlock("Dear hiring team,", "<w:p/>", ""),
        extract.ExtractedBlock("A paragraph about the role.", "<w:p/>", ""),
        extract.ExtractedBlock("Sincerely,", "<w:p/>", ""),
    ]
    assert roles.guess_roles(
        [b.text for b in blocks], [b.hint for b in blocks], "letter"
    ) == ["name", "salutation", "body_text", "closing"]


def test_prototype_falls_back_to_a_close_role():
    blueprint = store.blueprint_from_blocks(
        "cv",
        [extract.ExtractedBlock("A paragraph", "<w:p>proto</w:p>", "")],
        ["body_text"],
    )
    assert roles.prototype_for(blueprint, "bullet") == "<w:p>proto</w:p>"
    assert roles.prototype_for(blueprint, "no_such_role") is None


# --- the default --------------------------------------------------------------


def test_the_default_cv_covers_every_content_role():
    docx_bytes, blueprint = default.build("cv")
    assert docx_bytes
    present = set(blueprint.roles())
    for role in ("name", "headline", "contact", "section_title", "entry_title", "bullet", "skill_line"):
        assert role in present


def test_the_default_letter_covers_its_roles():
    _, blueprint = default.build("letter")
    assert {"name", "date", "salutation", "body_text", "closing", "signature"} <= set(
        blueprint.roles()
    )


# --- rendering ----------------------------------------------------------------


def test_rendering_replaces_the_sample_text_and_keeps_the_style():
    base, blueprint = default.build("cv")
    lines = [
        DraftLine(role="name", text="Camille Moreau"),
        DraftLine(role="headline", text="Senior Backend Engineer"),
        DraftLine(role="section_title", text="Experience"),
        DraftLine(role="bullet", text="Cut ingestion latency by 60%"),
    ]
    result = render.render(base, blueprint, lines)

    document = Document(BytesIO(result))
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    assert texts == [
        "Camille Moreau",
        "Senior Backend Engineer",
        "Experience",
        "Cut ingestion latency by 60%",
    ]
    # The sample content is gone, and the bullet keeps its list style.
    assert "Your Name" not in " ".join(texts)
    bullet = next(p for p in document.paragraphs if p.text.startswith("Cut ingestion"))
    assert bullet.style is not None and bullet.style.name == "List Bullet"
    name = next(p for p in document.paragraphs if p.text == "Camille Moreau")
    assert name.runs[0].bold is True
    assert name.runs[0].font.size == Pt(22)


def test_a_skills_line_keeps_its_category_style():
    document = Document()
    paragraph = document.add_paragraph()
    category = paragraph.add_run("Security:")
    category.bold = True
    paragraph.add_run(" OSCP · Audits")
    buffer = BytesIO()
    document.save(buffer)
    docx_bytes = buffer.getvalue()

    blueprint = store.blueprint_from_blocks(
        "cv", extract.extract_blocks(docx_bytes), ["skill_line"]
    )
    result = render.render(
        docx_bytes, blueprint, [DraftLine(role="skill_line", text="Languages: Rust, Python")]
    )

    rendered = Document(BytesIO(result)).paragraphs[0]
    assert rendered.text == "Languages: Rust, Python"
    assert [run.text for run in rendered.runs] == ["Languages:", " Rust, Python"]
    assert rendered.runs[0].bold is True
    assert rendered.runs[1].bold is not True


def test_a_skills_line_without_a_category_style_stays_one_run():
    base, blueprint = default.build("cv")
    rendered = Document(
        BytesIO(
            render.render(base, blueprint, [DraftLine(role="skill_line", text="Security: OSCP")])
        )
    ).paragraphs[0]
    assert rendered.text == "Security: OSCP"
    assert len(rendered.runs) == 1


def test_rendering_keeps_a_hyperlink_whose_text_is_still_there():
    docx_bytes = _linked_docx()
    blueprint = store.blueprint_from_blocks(
        "cv", extract.extract_blocks(docx_bytes), ["contact"]
    )
    result = render.render(
        docx_bytes,
        blueprint,
        [DraftLine(role="contact", text="camille@example.com · github.com/camille")],
    )

    paragraph = Document(BytesIO(result)).paragraphs[0]
    assert paragraph.text == "camille@example.com · github.com/camille"
    assert [(link.text, link.address) for link in paragraph.hyperlinks] == [
        ("camille@example.com", "mailto:camille@example.com")
    ]


def test_rendering_writes_a_link_as_text_when_its_text_is_gone():
    docx_bytes = _linked_docx()
    blueprint = store.blueprint_from_blocks(
        "cv", extract.extract_blocks(docx_bytes), ["contact"]
    )
    result = render.render(
        docx_bytes, blueprint, [DraftLine(role="contact", text="Write to me on LinkedIn")]
    )

    paragraph = Document(BytesIO(result)).paragraphs[0]
    assert paragraph.text == "Write to me on LinkedIn"
    assert paragraph.hyperlinks == []


def test_rendering_writes_link_text_as_plain_text_around_it():
    docx_bytes = _linked_docx()
    blueprint = store.blueprint_from_blocks(
        "cv", extract.extract_blocks(docx_bytes), ["contact"]
    )
    result = render.render(
        docx_bytes,
        blueprint,
        [DraftLine(role="contact", text="camille@example.com · github.com/camille")],
    )

    paragraph = Document(BytesIO(result)).paragraphs[0]
    # ``runs`` are the paragraph's own runs; a link is not one of them.
    assert "".join(run.text for run in paragraph.runs) == " · github.com/camille"


def test_rendering_creates_a_missing_section_from_a_close_prototype():
    base, complete = default.build("cv")
    # A template that styles section titles but has no bullet of its own.
    blueprint = TemplateBlueprint(
        kind="cv",
        blocks=[
            TemplateBlock(
                role="section_title",
                text="Section",
                xml=complete.prototype("section_title") or "",
            )
        ],
    )
    lines = [
        DraftLine(role="section_title", text="Experience"),
        DraftLine(role="bullet", text="A result"),
    ]
    document = Document(BytesIO(render.render(base, blueprint, lines)))
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    assert texts == ["Experience", "A result"]


def test_rendering_keeps_the_page_setup_of_the_base():
    base = Document(BytesIO(default.build("cv")[0]))
    base.sections[0].left_margin = Pt(99)
    buffer = BytesIO()
    base.save(buffer)

    document = Document(
        BytesIO(
            render.render(
                buffer.getvalue(),
                default.build("cv")[1],
                [DraftLine(role="name", text="Camille")],
            )
        )
    )
    assert document.sections[0].left_margin == Pt(99)


def test_rendering_escapes_nothing_and_keeps_spaces():
    base, blueprint = default.build("cv")
    document = Document(
        BytesIO(render.render(base, blueprint, [DraftLine(role="body_text", text="spaces  inside & <tags>")]))
    )
    assert document.paragraphs[0].text == "spaces  inside & <tags>"


# --- page counting ------------------------------------------------------------


def test_estimate_pages():
    assert pdf.estimate_pages([]) == 1
    assert pdf.estimate_pages([DraftLine()] * 46) == 2


def test_page_count_reports_an_estimate_without_a_converter(monkeypatch):
    monkeypatch.setattr("app.templates_engine.pdf.converter", lambda: None)
    pages, exact = pdf.page_count(b"", [DraftLine()] * 90)
    assert (pages, exact) == (2, False)


def test_page_count_uses_the_converter_when_it_is_there(monkeypatch):
    monkeypatch.setattr("app.templates_engine.pdf.count_pages", lambda *a, **k: 3)
    pages, exact = pdf.page_count(b"", [DraftLine()])
    assert (pages, exact) == (3, True)


# --- the store ----------------------------------------------------------------


def test_the_default_is_used_until_a_template_is_imported():
    assert store.has_custom("cv") is False
    blueprint = store.load_blueprint("cv")
    assert blueprint.kind == "cv"
    assert store.base_document("cv")


def test_an_imported_template_round_trips():
    _, blueprint = default.build("cv")
    store.save_custom("cv", default.build("cv")[0], blueprint)

    assert store.has_custom("cv") is True
    assert store.load_blueprint("cv").roles() == blueprint.roles()
    assert store.custom_blocks("cv")

    store.delete_custom("cv")
    assert store.has_custom("cv") is False


def test_corrected_roles_are_stored():
    docx_bytes = _docx_bytes(("Camille", {"bold": True}), ("A line", {}))
    store.save_custom("cv", docx_bytes, store.blueprint_from_blocks("cv", extract.extract_blocks(docx_bytes)))

    blueprint = store.save_roles("cv", ["headline", "not a role"])

    assert blueprint.roles() == ["headline", "body_text"]
    store.delete_custom("cv")
