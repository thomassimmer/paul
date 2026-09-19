"""A clean, single-column default, used when no template was provided.

The sample document and its blueprint are built together, so the prototypes are
known by construction: the default needs no analysis and no model call.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor

from app.models import TemplateBlock, TemplateBlueprint

ACCENT = RGBColor(0x2F, 0x6F, 0x4F)
MUTED = RGBColor(0x5B, 0x64, 0x70)


def build(kind: str) -> tuple[bytes, TemplateBlueprint]:
    """Return the default base document and the blueprint that goes with it."""
    builder = _build_letter if kind == "letter" else _build_cv
    return builder()


class _Builder:
    """Adds paragraphs to a sample document and keeps each one as a prototype."""

    def __init__(self, kind: str) -> None:
        self.document = Document()
        self.kind = kind
        self.blocks: list[TemplateBlock] = []
        section = self.document.sections[0]
        section.left_margin = section.right_margin = Pt(56)
        section.top_margin = section.bottom_margin = Pt(48)

    def add(
        self,
        role: str,
        text: str,
        *,
        size: float | None = None,
        bold: bool = False,
        italic: bool = False,
        color: RGBColor | None = None,
        list_item: bool = False,
        space_after: float = 4,
    ) -> None:
        paragraph = self.document.add_paragraph(style="List Bullet" if list_item else None)
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        paragraph.paragraph_format.space_after = Pt(space_after)
        self.blocks.append(
            TemplateBlock(role=role, text=text, xml=paragraph._p.xml)
        )

    def finish(self) -> tuple[bytes, TemplateBlueprint]:
        buffer = BytesIO()
        self.document.save(buffer)
        return buffer.getvalue(), TemplateBlueprint(kind=self.kind, blocks=self.blocks)


def _build_cv() -> tuple[bytes, TemplateBlueprint]:
    builder = _Builder("cv")
    builder.add("name", "Your Name", size=22, bold=True, space_after=2)
    builder.add("headline", "Your headline", size=11, color=ACCENT, space_after=2)
    builder.add("contact", "City · you@example.com · +00 000 00 00 00", size=9, color=MUTED)
    builder.add("section_title", "Experience", size=11, bold=True, color=ACCENT, space_after=2)
    builder.add("entry_title", "Job title — Company", bold=True, space_after=1)
    builder.add("entry_subtitle", "What the company does, and your remit there.", italic=True, space_after=1)
    builder.add("entry_dates", "2022-03 / 2024-06", size=9, color=MUTED)
    builder.add("bullet", "A result, with a number when you have one.", list_item=True)
    builder.add("body_text", "A short paragraph when a bullet is not enough.")
    builder.add("section_title", "Skills", size=11, bold=True, color=ACCENT, space_after=2)
    builder.add("skill_line", "Languages: Rust, Python")
    return builder.finish()


def _build_letter() -> tuple[bytes, TemplateBlueprint]:
    builder = _Builder("letter")
    builder.add("name", "Your Name", size=16, bold=True, space_after=2)
    builder.add("contact", "City · you@example.com · +00 000 00 00 00", size=9, color=MUTED)
    builder.add("date", "1 January 2026", size=9, color=MUTED, space_after=12)
    builder.add("recipient", "Hiring team, Company\nCity", size=10, space_after=12)
    builder.add("salutation", "Dear hiring team,", space_after=8)
    builder.add("body_text", "The paragraph that says why this role, in your own words.", space_after=8)
    builder.add("body_text", "A second paragraph, grounded in what you actually did.", space_after=8)
    builder.add(
        "closing",
        "Thank you for your time.",
        space_after=12,
    )
    builder.add("signature", "Your Name", bold=True)
    return builder.finish()
