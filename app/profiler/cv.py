"""CV text extraction.

Input is a PDF or a DOCX, as the README promises, plus Markdown/text so a CV
pasted from anywhere can be reused. Extraction is deterministic and is the only
thing that touches the uploaded bytes; the text is then treated as untrusted
data by the prompts.
"""

from __future__ import annotations

import io
from pathlib import Path

from app.profiler.text import tidy_text

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".txt", ".text"}
# Below this, the file is almost certainly a scan or an unsupported layout.
MIN_USEFUL_CHARS = 50


class CvError(Exception):
    """The CV could not be turned into text, for a reason we can explain."""


def extract_text(filename: str, data: bytes) -> str:
    """Return the text of a CV file, dispatching on its extension."""
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        text = _from_pdf(data)
    elif suffix == ".docx":
        text = _from_docx(data)
    elif suffix in {".md", ".txt", ".text"}:
        text = _decode(data)
    else:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise CvError(
            f"Unsupported file type “{suffix or filename}”. Use one of: {supported}. "
            "A Word .doc has to be saved as .docx first."
        )

    text = tidy_text(text)
    if len(text) < MIN_USEFUL_CHARS:
        raise CvError(
            "Almost no text could be extracted from this file. If it is a scan or an "
            "image-based PDF, paste the text instead."
        )
    return text


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise CvError("Could not decode the file as text.")


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on the install
        raise CvError(f"PDF support is not installed: {exc}") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # Many PDFs are "encrypted" with an empty password; try that, then give up.
            if not reader.decrypt(""):
                raise CvError("This PDF is password-protected: remove the password first.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except CvError:
        raise
    except Exception as exc:
        raise CvError(f"Could not read the PDF: {exc}") from exc


def _from_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - depends on the install
        raise CvError(f"DOCX support is not installed: {exc}") from exc

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise CvError(f"Could not read the DOCX: {exc}") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)
